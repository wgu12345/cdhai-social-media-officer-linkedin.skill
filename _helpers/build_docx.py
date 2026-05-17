#!/usr/bin/env python3
"""
Build the LinkedIn post .docx with inline embedded images, OR build the
merged before_you_post.docx report.

Two modes, dispatched by the --mode argument.

Usage:

    python3 build_docx.py --mode post --input <draft.md> --output <linkedin_post.docx> --manifest <image_manifest.json>

    python3 build_docx.py --mode report --input <report_data.json> --output <before_you_post.docx>

The manifest (for post mode) maps paragraph indices to image file paths:

    {
      "post_text": "full markdown of the post",
      "image_assignments": [
        {"paragraph_index": 1, "image_path": "/path/to/01.jpg", "caption": "Harang Ju keynote"},
        {"paragraph_index": 3, "image_path": "/path/to/05.jpg", "caption": "Equity panel"}
      ]
    }

Paragraph indices are 0-based, counting non-empty paragraphs in the post
(headlines and hashtags excluded — those are formatted differently).
"""
import argparse
import json
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("Missing python-docx. Run: pip install --break-system-packages python-docx")
    raise


SEVERITY_COLORS = {
    "CRITICAL": "C00000",
    "WARNING": "BF6900",
    "SUGGESTION": "555555",
    "PASS": "2E7D32",
}


def add_run(paragraph, text, *, bold=False, italic=False, color=None, size=None):
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    return run


def add_paragraph(doc, text, *, bold=False, italic=False, color=None, size=11, alignment=None):
    p = doc.add_paragraph()
    if alignment is not None:
        p.alignment = alignment
    add_run(p, text, bold=bold, italic=italic, color=color, size=size)
    return p


def add_heading(doc, text, level=1):
    if level == 1:
        p = doc.add_paragraph()
        add_run(p, text, bold=True, size=18, color="1F3864")
    elif level == 2:
        p = doc.add_paragraph()
        add_run(p, text, bold=True, size=14, color="2E75B6")
    else:
        p = doc.add_paragraph()
        add_run(p, text, bold=True, size=12)
    return p


def add_divider(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "2E75B6")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_image_inline(doc, image_path, caption=None, max_width_in=5.5):
    try:
        doc.add_picture(str(image_path), width=Inches(max_width_in))
        last_p = doc.paragraphs[-1]
        last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            cap_p = doc.add_paragraph()
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_run(cap_p, f"Figure: {caption}", italic=True, color="595959", size=9)
    except Exception as e:
        add_paragraph(doc, f"[Could not embed image: {image_path} — {e}]",
                      italic=True, color="C00000", size=9)


# ============================================================
# MODE: POST — LinkedIn post with inline images
# ============================================================

def split_post_text(post_text: str) -> dict:
    """Parse the post markdown into structured parts.

    Returns:
      {
        "title": "...",
        "paragraphs": ["...", "..."],   # body paragraphs in order
        "hashtags": "#A #B #C",
      }
    """
    lines = post_text.strip().split("\n")
    title = None
    paragraphs = []
    hashtags = None

    buf = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            if buf:
                paragraphs.append(" ".join(buf).strip())
                buf = []
            continue
        if title is None and stripped.startswith("# "):
            title = stripped[2:].strip()
            continue
        # Hashtag line (a paragraph that is only hashtags)
        if stripped.startswith("#") and all(
            tok.startswith("#") for tok in stripped.split()
        ):
            if buf:
                paragraphs.append(" ".join(buf).strip())
                buf = []
            hashtags = stripped
            continue
        buf.append(stripped)
    if buf:
        paragraphs.append(" ".join(buf).strip())

    return {"title": title, "paragraphs": paragraphs, "hashtags": hashtags}


def build_post(input_path: Path, output_path: Path, manifest_path: Path):
    with open(input_path, encoding="utf-8") as f:
        post_md = f.read()

    manifest = {}
    if manifest_path and manifest_path.exists():
        with open(manifest_path) as f:
            manifest = json.load(f)
    image_assignments = manifest.get("image_assignments", [])
    images_by_para = {a["paragraph_index"]: a for a in image_assignments}

    parsed = split_post_text(post_md)

    doc = Document()
    # Margins
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    # Title
    add_heading(doc, parsed["title"] or "LinkedIn Post Draft", level=1)
    add_divider(doc)

    # Body paragraphs with inline images
    for i, para in enumerate(parsed["paragraphs"]):
        p = doc.add_paragraph()
        add_run(p, para, size=12)
        if i in images_by_para:
            img = images_by_para[i]
            add_image_inline(doc, img["image_path"], caption=img.get("caption"))

    # Hashtags
    if parsed["hashtags"]:
        add_divider(doc)
        p = doc.add_paragraph()
        add_run(p, parsed["hashtags"], color="2E75B6", size=11)

    doc.save(str(output_path))


# ============================================================
# MODE: REPORT — before_you_post.docx merged report
# ============================================================

def build_report(input_path: Path, output_path: Path):
    """Build before_you_post.docx from a structured JSON.

    Expected input schema (omit any section if empty):

    {
      "run_meta": {
        "folder": "...", "skill_version": "0.4.0", "date": "2026-05-16",
        "post_type": "event_recap", "draft_chars": 1284
      },
      "must_verify": [
        {"severity": "CRITICAL"|"WARNING", "text": "..."}
      ],
      "decisions": ["string", ...],
      "image_picks": [
        {"image": "01.jpg", "key_point": "...", "source": "user-provided",
         "match_confidence": 0.88, "rationale": "..."}
      ],
      "quality_checks": [
        {"name": "key_points_coverage", "status": "PASS"|"FLAG", "detail": "..."},
        ...
      ],
      "ai_generated_disclosures": [
        {"image": "fallback_x.jpg", "prompt": "...", "tool": "dalle"}
      ],
      "date_verification": {
        "today": "2026-05-16", "event_date": "2026-04-28", "computed_phrase": "earlier this month",
        "conflicts": ["..."]
      },
      "paper_info": {  // only for research_published runs
        "title": "...", "authors": [...], "venue": "...", "doi": "..."
      },
      "posting_instructions": ["step 1", "step 2", ...]
    }
    """
    with open(input_path, encoding="utf-8") as f:
        data = json.load(f)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    # Title
    add_heading(doc, "Before You Post — linkedin_post.docx", level=1)

    # ---- LEGAL WARNING for web-sourced images — must be FIRST, before everything else ----
    web_sourced = data.get("web_sourced_disclosures", [])
    if web_sourced:
        # Big red warning box at the top
        warn_p = doc.add_paragraph()
        pPr = warn_p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        for side in ("top", "left", "bottom", "right"):
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "12")
            b.set(qn("w:color"), "C00000")
            pBdr.append(b)
        pPr.append(pBdr)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "FFF2F2")
        pPr.append(shd)
        add_run(warn_p, "⚠ WEB-SOURCED IMAGES — VERIFY USAGE RIGHTS BEFORE POSTING",
                bold=True, color="C00000", size=14)

        bp = doc.add_paragraph()
        add_run(bp,
                f"This post contains {len(web_sourced)} image(s) sourced from web search. ",
                bold=True, color="C00000", size=11)
        add_run(bp,
                "Each must be cleared for use on a JHU / CDHAI / Carey channel "
                "before publishing. Acceptable clearances: public-domain status "
                "verified at the source; explicit written permission from the rights "
                "holder; or a license that permits the intended use (CC-BY, CC0, etc.). "
                "Random web images are typically copyrighted; posting without rights "
                "is institutional legal exposure.",
                size=11)

        for w in web_sourced:
            wp = doc.add_paragraph(style="List Bullet")
            add_run(wp, f"{w.get('image', 'unknown')}", bold=True, size=11)
            sp = doc.add_paragraph()
            sp.paragraph_format.left_indent = Inches(0.5)
            add_run(sp, f"  Source page: ", italic=True, color="595959", size=10)
            add_run(sp, w.get("source_page", "—"), italic=True, color="0563C1", size=10)
            sp2 = doc.add_paragraph()
            sp2.paragraph_format.left_indent = Inches(0.5)
            add_run(sp2, f"  Search query: \"{w.get('search_query', '—')}\"  ·  Title: \"{w.get('title', '—')}\"",
                    italic=True, color="595959", size=10)
            sp3 = doc.add_paragraph()
            sp3.paragraph_format.left_indent = Inches(0.5)
            add_run(sp3, "  ☐ Rights verified before posting (check box manually)",
                    italic=True, color="C00000", size=10)

        add_divider(doc)

    # Run meta
    meta = data.get("run_meta", {})
    if meta:
        p = doc.add_paragraph()
        add_run(p, f"Folder: ", bold=True, size=10)
        add_run(p, f"{meta.get('folder', 'unknown')}    ", size=10)
        add_run(p, f"Skill version: ", bold=True, size=10)
        add_run(p, f"{meta.get('skill_version', '0.4.0')}    ", size=10)
        add_run(p, f"Date: ", bold=True, size=10)
        add_run(p, f"{meta.get('date', '')}    ", size=10)
        add_run(p, f"Post type: ", bold=True, size=10)
        add_run(p, f"{meta.get('post_type', '')}", size=10)

    add_divider(doc)

    # ---- Must verify section ----
    must_verify = data.get("must_verify", [])
    add_heading(doc, f"⚠️ Must verify ({len(must_verify)} items)", level=2)
    if not must_verify:
        add_paragraph(doc, "No critical items. Ready to post after reading the rest of this report.",
                      italic=True, color="2E7D32")
    else:
        for item in must_verify:
            sev = item.get("severity", "WARNING")
            color = SEVERITY_COLORS.get(sev, "BF6900")
            p = doc.add_paragraph(style="List Bullet")
            add_run(p, f"{sev}: ", bold=True, color=color, size=11)
            add_run(p, item.get("text", ""), size=11)

    # ---- Date verification ----
    dv = data.get("date_verification")
    if dv:
        add_heading(doc, "📅 Date verification", level=2)
        p = doc.add_paragraph()
        add_run(p, "Today: ", bold=True); add_run(p, f"{dv.get('today', '')}    ")
        add_run(p, "Event date: ", bold=True); add_run(p, f"{dv.get('event_date', '')}    ")
        add_run(p, "Phrasing used in post: ", bold=True); add_run(p, f"{dv.get('computed_phrase', '')}")
        for conflict in dv.get("conflicts", []) or []:
            p = doc.add_paragraph(style="List Bullet")
            add_run(p, "Conflict: ", bold=True, color="BF6900")
            add_run(p, conflict)

    # ---- Decisions ----
    decisions = data.get("decisions", [])
    if decisions:
        add_heading(doc, "✓ Decisions you should know", level=2)
        for d in decisions:
            p = doc.add_paragraph(style="List Bullet")
            add_run(p, d, size=11)

    # ---- Image picks ----
    image_picks = data.get("image_picks", [])
    if image_picks:
        add_heading(doc, f"🖼 Image picks ({len(image_picks)} matched)", level=2)
        for pick in image_picks:
            p = doc.add_paragraph(style="List Bullet")
            add_run(p, f"{pick.get('image', 'unknown')}: ", bold=True, size=11)
            add_run(p, f"{pick.get('rationale', '')}", size=11)
            sub = doc.add_paragraph()
            sub.paragraph_format.left_indent = Inches(0.5)
            add_run(sub, f"  Key Point: {pick.get('key_point', '—')}", italic=True, color="595959", size=10)
            sub2 = doc.add_paragraph()
            sub2.paragraph_format.left_indent = Inches(0.5)
            confidence = pick.get('match_confidence')
            subject = pick.get('subject', '—')
            scene = pick.get('scene_type', '—')
            add_run(sub2,
                    f"  Source: {pick.get('source', '—')}  ·  Match confidence: {confidence if confidence is not None else 'n/a'}  ·  Subject: {subject}  ·  Scene: {scene}",
                    italic=True, color="595959", size=10)

    # ---- Diversity check ----
    diversity = data.get("diversity_check")
    if diversity:
        add_heading(doc, "🎨 Diversity check", level=2)
        status = diversity.get("status", "PASS")
        scolor = SEVERITY_COLORS.get(status, "595555")
        p = doc.add_paragraph()
        add_run(p, "Status: ", bold=True)
        add_run(p, status, bold=True, color=scolor)
        if diversity.get("subjects"):
            p = doc.add_paragraph()
            add_run(p, "Subject distribution: ", bold=True, size=10)
            sub_str = ", ".join([f"{name} × {count}" for name, count in diversity["subjects"].items()])
            add_run(p, sub_str, size=10, color="595959")
        if diversity.get("scenes"):
            p = doc.add_paragraph()
            add_run(p, "Scene distribution: ", bold=True, size=10)
            scene_str = ", ".join([f"{name} × {count}" for name, count in diversity["scenes"].items()])
            add_run(p, scene_str, size=10, color="595959")
        if diversity.get("notes"):
            p = doc.add_paragraph()
            add_run(p, diversity["notes"], italic=True, color="595959", size=10)

    # ---- Quality checks ----
    qc = data.get("quality_checks", [])
    if qc:
        add_heading(doc, "✓ Quality checks", level=2)
        for check in qc:
            name = check.get("name", "")
            status = check.get("status", "")
            color = SEVERITY_COLORS.get(status, "595555")
            p = doc.add_paragraph(style="List Bullet")
            add_run(p, f"{name}: ", bold=True, size=11)
            add_run(p, status, bold=True, color=color, size=11)
            detail = check.get("detail", "")
            if detail:
                add_run(p, f"  —  {detail}", size=11, color="595959")

    # ---- Paper info (for research_published) ----
    paper = data.get("paper_info")
    if paper:
        add_heading(doc, "📄 Paper details", level=2)
        p = doc.add_paragraph()
        add_run(p, "Title: ", bold=True); add_run(p, paper.get("title", ""))
        if paper.get("authors"):
            p = doc.add_paragraph()
            add_run(p, "Authors: ", bold=True)
            author_str = ", ".join(
                a.get("name", "") if isinstance(a, dict) else str(a)
                for a in paper["authors"]
            )
            add_run(p, author_str)
        if paper.get("venue"):
            p = doc.add_paragraph()
            add_run(p, "Venue: ", bold=True); add_run(p, paper["venue"])
        if paper.get("doi"):
            p = doc.add_paragraph()
            add_run(p, "DOI: ", bold=True); add_run(p, paper["doi"])

    # ---- AI-generated disclosures ----
    ai_disc = data.get("ai_generated_disclosures", [])
    if ai_disc:
        add_heading(doc, "🤖 AI-generated images used (disclosed per JHU policy)", level=2)
        for disc in ai_disc:
            p = doc.add_paragraph(style="List Bullet")
            add_run(p, f"{disc.get('image', 'unknown')}", bold=True, size=10)
            add_run(p, f"  ·  Tool: {disc.get('tool', '')}", size=10)
            add_run(p, f"  ·  Prompt: \"{disc.get('prompt', '')}\"", size=10, italic=True, color="595959")

    # ---- Posting instructions ----
    instructions = data.get("posting_instructions", [])
    if instructions:
        add_heading(doc, "📤 To publish", level=2)
        for i, step in enumerate(instructions, 1):
            p = doc.add_paragraph(style="List Number")
            add_run(p, step, size=11)

    doc.save(str(output_path))


# ============================================================
# CLI
# ============================================================

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=["post", "report"], required=True)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--manifest", type=Path, help="image manifest (post mode only)")
    args = parser.parse_args()

    if args.mode == "post":
        build_post(args.input, args.output, args.manifest)
    else:
        build_report(args.input, args.output)
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
