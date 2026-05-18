#!/usr/bin/env python3
"""
validate_template.py — v0.4.4 (soft mode)

Diagnostic checker for the content template. v0.4.4 switched this from
a HARD gate (exit 2 on missing fields) to a SOFT diagnostic: it always
exits 0 when the folder is readable, and reports findings in a JSON
file that Phase 9 surfaces in `before_you_post.docx` under the Audit
Log and Must-Verify sections.

Rationale: the v0.4.3 hard gate produced false positives in real use
(legitimate "NA" / "no special notes" answers got blocked). The team
trade-off: warn loudly in the report instead of refusing to run, and
keep the reviewer gate (reviewer_check.py) hard so quality is still
protected.

Exit codes:
    0  → diagnostic completed (regardless of whether fields are missing)
    2  → genuine error: folder unreadable or argv malformed
    3  → environment problem (python-docx not installed)

Always writes a JSON report to:
    ~/.cdhai-linkedin-skill/cache/template_validation.json

Usage:
    python3 _helpers/validate_template.py <folder_path>
"""

import json
import re
import sys
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("ERROR: python-docx not installed. Run: pip3 install python-docx",
          file=sys.stderr)
    sys.exit(3)

REQUIRED_FIELDS = [
    "Purpose",
    "Post Type",
    "Date",
    "Key Points",
    "Background",
    "Notes",
]

# Each required field can appear under multiple label variants in the
# user's template. We try the longest variant first so we match the
# whole label and don't mistake its tail for body content.
FIELD_LABEL_VARIANTS = {
    "Purpose":    [r"Purpose"],
    "Post Type":  [r"Post Type", r"Post\s*Type"],
    "Date":       [r"Date"],
    "Key Points": [r"Key Points", r"Key\s*Points"],
    "Background": [r"Background\s*/\s*Source\s*Materials",
                   r"Background\s+and\s+Source\s+Materials",
                   r"Background"],
    "Notes":      [r"Notes\s+for\s+the\s+writer",
                   r"Writer\s*Notes",
                   r"Notes"],
}

# A field counts as "filled enough" if it has this many real characters
# after the label. NA / N/A / None on their own count as sparse.
MIN_CONTENT_CHARS = 10
SPARSE_PLACEHOLDERS = {"na", "n/a", "none", "tbd", "tbc", "-", "—", ".", ""}

CACHE_DIR = Path.home() / ".cdhai-linkedin-skill" / "cache"
REPORT_PATH = CACHE_DIR / "template_validation.json"


def find_content_docx(folder: Path):
    """Return the most plausible content template, or None."""
    candidates = list(folder.rglob("*.docx"))
    if not candidates:
        return None
    for c in candidates:
        if c.name.lower() == "content.docx":
            return c
    for c in candidates:
        if "template" in c.stem.lower() or "content" in c.stem.lower():
            return c
    return candidates[0]


def extract_all_text(doc_path: Path) -> str:
    doc = Document(doc_path)
    chunks = []
    for p in doc.paragraphs:
        chunks.append(p.text)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    return "\n".join(chunks)


def classify_field(text: str, field: str) -> str:
    """Return one of: 'filled', 'sparse', 'missing'."""
    variants = FIELD_LABEL_VARIANTS.get(field, [re.escape(field)])

    # Try each variant in order (longest / most specific first).
    match = None
    for v in variants:
        pattern = re.compile(
            rf"(^|\n)\s*({v})\s*[:\-]\s*(.*)",
            re.IGNORECASE | re.MULTILINE,
        )
        m = pattern.search(text)
        if m:
            match = m
            break

    if not match:
        return "missing"

    same_line = match.group(3).strip()
    after = text[match.end(): match.end() + 400].strip()
    body = (same_line + " " + after).strip()

    # Strip body up to the next field label of any kind.
    all_variants = [v for vs in FIELD_LABEL_VARIANTS.values() for v in vs]
    earliest = len(body)
    for v in all_variants:
        nxt = re.search(rf"\b{v}\b\s*[:\-]", body, re.IGNORECASE)
        if nxt and nxt.start() < earliest:
            earliest = nxt.start()
    body = body[:earliest].strip()

    plain = body.lower().strip().rstrip(".:;,")
    if plain in SPARSE_PLACEHOLDERS:
        return "sparse"
    if len(body) < MIN_CONTENT_CHARS:
        return "sparse"
    return "filled"


def write_report(report: dict):
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text(json.dumps(report, indent=2), encoding="utf-8")


def main():
    if len(sys.argv) < 2:
        print("Usage: validate_template.py <folder_path>", file=sys.stderr)
        sys.exit(2)

    folder = Path(sys.argv[1]).expanduser().resolve()
    if not folder.is_dir():
        print(f"ERROR: folder not found: {folder}", file=sys.stderr)
        sys.exit(2)

    timestamp = datetime.now().isoformat(timespec="seconds")
    doc_path = find_content_docx(folder)

    # Case 1: no docx at all → WARN but don't block
    if doc_path is None:
        report = {
            "status": "NO_TEMPLATE",
            "timestamp": timestamp,
            "template_path": None,
            "missing_fields": REQUIRED_FIELDS,
            "sparse_fields": [],
            "filled_fields": [],
            "summary_for_audit_log":
                "Template validation: NO TEMPLATE FOUND — skill proceeded "
                "with prompt-only input. All six standard fields are "
                "missing; treat output as best-effort.",
            "summary_for_must_verify": (
                "No content.docx template was found in this folder. The "
                "draft was generated from chat prompts and supporting "
                "materials only. Manually verify every factual claim "
                "before posting."
            ),
        }
        write_report(report)
        print(
            "WARN: no .docx template found in the folder. Continuing in "
            "soft mode (v0.4.4). The before_you_post.docx report will "
            "flag this prominently.\n"
            f"Report written to {REPORT_PATH}",
            file=sys.stderr,
        )
        sys.exit(0)

    # Case 2: docx exists — classify each field
    text = extract_all_text(doc_path)
    missing, sparse, filled = [], [], []
    for f in REQUIRED_FIELDS:
        c = classify_field(text, f)
        if c == "missing":
            missing.append(f)
        elif c == "sparse":
            sparse.append(f)
        else:
            filled.append(f)

    if not missing and not sparse:
        report = {
            "status": "OK",
            "timestamp": timestamp,
            "template_path": str(doc_path),
            "missing_fields": [],
            "sparse_fields": [],
            "filled_fields": filled,
            "summary_for_audit_log":
                "Template validation: PASS (all six fields filled).",
            "summary_for_must_verify": None,
        }
        write_report(report)
        print(f"OK: template validated at {doc_path.name}")
        sys.exit(0)

    # Case 3: docx exists but some fields missing/sparse → WARN but proceed
    parts = []
    if missing:
        parts.append(f"missing: {', '.join(missing)}")
    if sparse:
        parts.append(f"sparse (NA / blank / <10 chars): {', '.join(sparse)}")
    summary_short = "; ".join(parts)

    report = {
        "status": "WARN",
        "timestamp": timestamp,
        "template_path": str(doc_path),
        "missing_fields": missing,
        "sparse_fields": sparse,
        "filled_fields": filled,
        "summary_for_audit_log":
            f"Template validation: WARN — {summary_short}. "
            "Skill proceeded; AI may have inferred or omitted content "
            "for these fields. Review carefully.",
        "summary_for_must_verify": (
            "Some template fields were missing or sparse. The skill "
            "did not stop, but you should verify what the AI filled in "
            "for these fields:\n"
            + "\n".join(f"  • {f}" for f in (missing + sparse))
        ),
    }
    write_report(report)
    print(
        f"WARN: template at {doc_path.name} has gaps — {summary_short}.\n"
        "Continuing in soft mode (v0.4.4). The before_you_post.docx "
        "report will flag this in the Audit Log and Must-Verify "
        "sections.\n"
        f"Report written to {REPORT_PATH}",
        file=sys.stderr,
    )
    sys.exit(0)


if __name__ == "__main__":
    main()
